"""Тесты сборщика отчётов: чтение CSV, роли колонок, порядок иерархии, значения вместо формул."""
import csv
import os
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import report_builder as R  # noqa: E402

HEAD = ["Уровень", "Наименование", "GROUP_NAME", "OTD_NAME", "Из них",
        "Случаи 2026", "Сумма 2026, млн руб", "Случаи 2025", "Сумма 2025, млн руб"]


def write_csv(path, header, rows, encoding="cp1251", sep=";"):
    with open(path, "w", encoding=encoding, newline="") as f:
        w = csv.writer(f, delimiter=sep)
        w.writerow(header)
        w.writerows(rows)


def build(src, **cfg_extra):
    cfg = R.deep_merge(R.DEFAULTS, cfg_extra)
    out = str(src) + ".xlsx"
    R.build(str(src), cfg, out)
    return load_workbook(out)


@pytest.fixture
def hierarchy_dir(tmp_path):
    """Корректная структура: дети под родителем, «из них» — сразу под своей строкой, соседи перемешаны."""
    rows = [[0, "Всего", "", "", "", 600, 60.0, 500, 50.0]]
    for g, s in (("Терапия", 10.0), ("Хирургия", 40.0), ("Офтальмология", 10.0)):
        rows.append([1, g, g, "", "", 100, s, 100, s / 2])
        rows.append([1, "из них операции", g, "", "операции", 10, 1.0, 10, 1.0])
        for o, so in ((0, 1.0), (1, 5.0), (2, 3.0)):
            rows.append([2, f"{g} {o}", g, f"{g} {o}", "", 10, so, 10, so])
    write_csv(tmp_path / "rep_1_5.csv", HEAD, rows)
    return tmp_path


def test_to_numbers_comma_spaces_percent():
    nums, all_pct = R.to_numbers(pd.Series(["1 234,50", "-10", "1,234.5"]))
    assert list(nums) == [1234.5, -10.0, 1234.5]
    assert not all_pct
    nums, all_pct = R.to_numbers(pd.Series(["12,5%", "3%"]))
    assert list(nums) == [0.125, 0.03] and all_pct


def test_fix_header_restores_delta():
    assert R.fix_header("? Случаи") == "Δ Случаи"
    assert R.fix_header("?% Сумма") == "Δ% Сумма"


def test_roles_leading_zero_codes_dates_and_metrics(tmp_path):
    write_csv(tmp_path / "t.csv", ["Регион", "Код", "Дата", "Выручка 2026", "Выручка 2025", "Доля"],
              [["Москва", "01", "01.10.2026", "10,5", "5", "12,5%"],
               ["Казань", "02", "02.10.2026", "3", "0", "1%"]], encoding="utf-8", sep=",")
    rep = R.prepare(str(tmp_path / "t.csv"), "t.csv", R.file_config(R.DEFAULTS, "t.csv"), "")
    by = {c.name: c for c in rep.columns}
    assert by["Код"].role == "dimension" and by["Код"].values == ["01", "02"]
    assert by["Дата"].kind == "date"
    assert by["Выручка 2026"].kind == "money" and by["Доля"].kind == "pct"
    assert by["Δ Выручка"].values == [5.5, 3.0]
    assert by["Δ% Выручка"].values == [1.1, None]          # при 2025 = 0 — пусто


def test_hierarchy_sorted_inside_groups_subrows_stay(hierarchy_dir):
    ws = build(hierarchy_dir).worksheets[1]
    names = [r[1] for r in ws.iter_rows(min_row=8, values_only=True) if r[0] is not None]
    assert names == ["Всего",
                     "Хирургия", "из них операции", "Хирургия 1", "Хирургия 2", "Хирургия 0",
                     "Терапия", "из них операции", "Терапия 1", "Терапия 2", "Терапия 0",
                     "Офтальмология", "из них операции", "Офтальмология 1", "Офтальмология 2", "Офтальмология 0"]


def test_values_only_and_formulas_sheet(hierarchy_dir):
    wb = build(hierarchy_dir)
    for ws in wb.worksheets:
        if ws.title == R.FORMULAS_SHEET:
            continue
        assert not any(c.data_type == "f" for row in ws.iter_rows() for c in row), ws.title
        assert len(ws.conditional_formatting) == 0
    fs = wb[R.FORMULAS_SHEET]
    assert fs.sheet_state == "hidden"
    texts = [r[3] for r in fs.iter_rows(min_row=6, values_only=True)]
    assert any(t.startswith("=IF(N(") for t in texts)
    assert all(fs.cell(6 + i, 4).data_type == "s" for i in range(len(texts)))


def test_live_mode_writes_formulas(hierarchy_dir):
    wb = build(hierarchy_dir, table={"formulas": "live"})
    ws = wb.worksheets[1]
    assert any(c.data_type == "f" for row in ws.iter_rows() for c in row)


def test_flat_totals_are_values(tmp_path):
    write_csv(tmp_path / "s.csv", ["Товар", "Кол-во 2026", "Кол-во 2025"], [["a", 5, 2], ["b", 1, 4]])
    ws = build(tmp_path).worksheets[1]
    total = [ws.cell(5, i).value for i in range(1, 6)]
    assert total == ["Итого", 6, 6, 0, 0.0]


def test_report_number_and_skip(tmp_path):
    assert R.report_number("rep_app_stac_1_1_202610021617") == ("1.1", "rep app stac")
    assert R.report_number("19.1 ЖК") == ("19.1", "ЖК")
    assert R.report_number("продажи") == ("", "продажи")
    write_csv(tmp_path / "a.csv", ["x", "n"], [["a", 1]])
    write_csv(tmp_path / "черновик.csv", ["x", "n"], [["a", 1]])
    wb = build(tmp_path, files={"*черновик*": {"skip": True}})
    assert len([w for w in wb.worksheets if w.title not in (R.TOC_SHEET, R.HELP_SHEET, R.FORMULAS_SHEET)]) == 1


def fill_of(cell):
    return cell.fill.fgColor.rgb[-6:] if cell.fill and cell.fill.fill_type == "solid" else None


def test_palette_b_default_and_a_switch(hierarchy_dir):
    ws = build(hierarchy_dir).worksheets[1]
    assert fill_of(ws.cell(8, 2)) == "C9DCEC"                 # «Всего» — уровень 0 палитры B
    assert fill_of(ws.cell(9, 2)) == "DCE8F3"                 # уровень 1
    assert fill_of(ws.cell(5, 2)) == "E8ECF0"                 # шапка
    assert ws.cell(2, 2).font.color.rgb[-6:] == "1B3A5C"      # заголовок — глубокий синий
    detail = next(r for r in range(8, 30) if ws.cell(r, 2).value == "Хирургия 1")
    assert fill_of(ws.cell(detail, 4)) == "ECEFF2"            # ключевая колонка «Сумма 2026»
    ws = build(hierarchy_dir, style={"palette": "A"}).worksheets[1]
    assert fill_of(ws.cell(8, 2)) == "F4B183"
    ws = build(hierarchy_dir, style={"colors": {"levels": {0: "112233"}}}).worksheets[1]
    assert fill_of(ws.cell(8, 2)) == "112233"
