"""
Универсальная сборка Excel-отчёта из папки CSV: каждый CSV -> отдельный лист в едином оформлении
(стиль — из build_app_excel12.py «Распределение амбулаторных услуг в стационарах») + лист «Содержание»
+ лист «Справка по колонкам».

Что делает сам, без настроек:
  • читает CSV: кодировка (UTF-8 / Windows-1251), разделитель (; , TAB |), десятичная запятая, «12,5%»,
    «\\N»/NULL — пусто; «? » в начале заголовка (Δ, потерянная при выгрузке в cp1251) -> «Δ »;
  • определяет роль каждой колонки:
      level      — «Уровень» (0 = Всего, 1, 2 …) -> иерархический отчёт: структура — как в CSV (родитель, под ним дети;
                   «из них …» — сразу под своей строкой), сортируются только дети внутри групп; отступы,
                   заливка итоговых уровней, жирный «Всего» и 1-й уровень (как в отчётах АПП);
      name       — «Наименование» (в иерархическом отчёте — главная колонка);
      dimension  — разрезы/текст/коды/даты (коды с ведущими нулями остаются текстом);
      metric     — числовые колонки; вид: count (целые), money (дробные / «сумма», «руб»), pct («%», «доля»),
                   delta («Δ», «дельта», «изм»), delta_pct («Δ%»);
  • колонки «X 2026» / «X 2025» -> шапка в 2 строки по годам; если Δ в CSV нет — добавляет Δ и Δ%;
    Δ% больше 200% показывает как «> 200%»;
  • в ячейках — только значения и заливки (файл не тормозит, ничего не «плывёт»); формулы (Δ, Δ%, Итого) и
    правила раскраски записаны текстом: в примечании к заголовку колонки и на скрытом листе «Формулы» —
    скопировать формулу в первую ячейку диапазона и протянуть вниз (table.formulas: live — формулы прямо в ячейках);
  • раскраска Δ и Δ% мягкой цветовой шкалой (рост — красный, снижение — зелёный; можно наоборот),
    по умолчанию — заливкой ячеек (без условного форматирования), можно — условным форматированием;
  • плоские таблицы: строка «Итого» над шапкой (формула SUBTOTAL по видимым строкам — на листе «Формулы»), автофильтр;
  • ширина колонок по содержимому, закреплённая шапка, печать по ширине, «Страница N из M», ссылки
    «← Содержание» на каждом листе.

Всё, что нужно вписать руками (названия, описания, подписи колонок, какие колонки метрики, что красить),
задаётся в конфиге report.yaml в папке с CSV. Шаблон с найденными файлами и колонками:
  python report_builder.py "D:\\папка_с_csv" --init       -> создаст report.yaml (существующий не трогает)
  python report_builder.py "D:\\папка_с_csv" --explain    -> покажет, как распознаны колонки, без сборки
  python report_builder.py "D:\\папка_с_csv"              -> соберёт <папка>.xlsx рядом с папкой
  python report_builder.py "D:\\папка" --out "D:\\итог.xlsx" --config "D:\\общий_конфиг.yaml"

Нужны: pip install pandas openpyxl pyyaml   (pyyaml — только для конфига .yaml; .json работает без него)
Цифры из CSV не пересчитываются: в ячейки попадают значения CSV; добавленные Δ / Δ% / Итого посчитаны скриптом
и записаны числами, их формулы — на листе «Формулы» и в примечаниях к заголовкам.
"""
import argparse
import copy
import csv
import datetime as dt
import fnmatch
import json
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

try:
    import yaml
except ImportError:
    yaml = None

CONFIG_NAMES = ("report.yaml", "report.yml", "report.json")
TOC_SHEET, HELP_SHEET, FORMULAS_SHEET = "Содержание", "Справка по колонкам", "Формулы"
EXCEL_MAX_ROWS = 1_048_576

# ---------------------------------------------------------------- настройки по умолчанию (переопределяются конфигом)
DEFAULTS = {
    "book": {
        "title": None,                 # заголовок книги; по умолчанию — имя папки
        "subtitle": "",                # красная строка под заголовком (на всех листах)
        "period": "",                  # подпись периода к годам в шапке: «2026 9м»
        "notes": [],                   # пояснения на листе «Содержание» (список строк)
        "output": None,                # путь к .xlsx; по умолчанию <папка>.xlsx рядом с папкой
        "columns_sheet": True,         # лист «Справка по колонкам»
        "formulas_sheet": "hidden",    # лист «Формулы» (формулы и правила раскраски текстом): hidden | visible | off
        "header_notes": True,          # примечание с формулой / правилом раскраски на заголовке колонки
    },
    "style": {
        "font": "Arial Narrow", "size": 14,
        "title_color": "002060", "subtitle_color": "C00000", "muted_color": "808080",
        "view": "page_break",          # page_break (страничный, как в АПП) | normal
        "zoom": 70,
        "paper": "auto",               # auto (A3, если таблица широкая) | A3 | A4
        "orientation": "landscape",
    },
    "table": {
        "compare": "auto",             # auto — добавить Δ / Δ% для пар «X 2026 / X 2025»; off — не добавлять
        "formulas": "values",          # values — в ячейках числа, формулы — текстом на листе «Формулы» | live
        "totals": "auto",              # плоские таблицы: auto/on — строка «Итого» над шапкой; off — нет
        "autofilter": True,            # плоские таблицы
        "sort": None,                  # «-Сумма 2026» — по убыванию; «Наименование» — по возрастанию
        "hierarchy_sort": "auto",      # иерархия: auto — дети по убыванию главной метрики; off — порядок CSV
        "row_numbers": "auto",         # колонка «№»: auto — только в иерархических отчётах
        "column_numbers": True,        # строка с номерами колонок под шапкой
        "freeze_columns": "auto",      # сколько колонок закрепить слева; auto — колонки до первой метрики (≤3)
        "max_width": 45, "name_max_width": 90,
    },
    "coloring": {
        "mode": "static",              # static — цвет заливкой (рекомендуется) | conditional — условное форм. | off
        "positive_is": "bad",          # bad — рост красный, снижение зелёный | good — наоборот
        "columns": None,               # список колонок для раскраски; None — все Δ и Δ%
        "gamma": 2.5,                  # насыщенность = доля^gamma: средние значения еле заметны
        "pct_cap": 2.0,                # Δ% больше 200% -> «> 200%», край шкалы Δ%
    },
    "columns": {},                     # правила по именам колонок (для всех файлов): см. --init
    "files": {},                       # настройки отдельных файлов по маске имени: см. --init
}

# ---------------------------------------------------------------- палитра (из build_app_excel12.py)
C_BLACK, C_GREY, C_GREY_DARK, C_GREY_LIGHT, C_SERVICE = "000000", "808080", "767171", "BFBFBF", "595959"
HEAD_FILLS = {"dimension": "F2F2F2", "count": "DEEBF7", "money": "E2F0D9", "number": "F2F2F2",
              "pct": "FFF2CC", "delta": "FFF2CC", "delta_pct": "FFF2CC", "service": "F2F2F2"}
GROUP_FILL_YEAR, GROUP_FILL_DELTA = "F2F2F2", "FFF2CC"
LEVEL_FILLS = {0: "F4B183", 1: "FBE5D6", 2: "FFF2CC", 3: "DDEBF7", 4: "EDEDED"}
TOTAL_FILL = "F4B183"
CF_RED, CF_GREEN, CF_WHITE = "FCE0E0", "C8E8D0", "FFFFFF"
LINK_COLOR = "0563C1"
HEAD_SIDE = Side(style="thin", color=C_GREY_LIGHT)
DATA_SIDE = Side(style="thin", color="D9D9D9")
HEAD_BORDER = Border(left=HEAD_SIDE, right=HEAD_SIDE, top=HEAD_SIDE, bottom=HEAD_SIDE)
DATA_BORDER = Border(left=DATA_SIDE, right=DATA_SIDE, top=DATA_SIDE, bottom=DATA_SIDE)

FMT = {
    "count": r"#,##0;\-#,##0;\-",
    "money": r"#,##0.00;\-#,##0.00;\-",
    "number": r"#,##0.##;\-#,##0.##;\-",
    "pct": r"0.0%",
    "pct_points": r'0.0"%"',
    "delta_pct": r"\+0.0%;\-0.0%;0.0%",
    "date": "DD.MM.YYYY",
    "datetime": "DD.MM.YYYY HH:MM",
    "text": "@",
}
FMT_PCT_CAP = '"> 200%"'
CHAR_W, HEAD_CHAR_W = 1.3, 1.45   # ширина символа (данные / жирная шапка) в единицах Excel при размере 14

# ---------------------------------------------------------------- распознавание по имени колонки
RX = {
    "level": re.compile(r"^(уровень|level|lvl)$", re.I),
    "name": re.compile(r"^(наименование|название|name)$", re.I),
    "subrow": re.compile(r"^из них$", re.I),
    "id": re.compile(r"(^|[_\s])(id|code|код|кода|инн|огрн|кпп|снилс|окпо|окато|октмо|gid|pat_id|полис|номер)"
                     r"($|[_\s])|^(№|n|no|nn|пп|№ п/п)$", re.I),
    "period": re.compile(r"^(год|year|месяц|month|квартал|quarter|период|period|дата|date)\b", re.I),
    "delta_pct": re.compile(r"^(Δ|δ)\s*%|^(дельта|delta|изм\w*)\s*%", re.I),
    "delta": re.compile(r"^(Δ|δ|дельта|delta|diff|разница|изм)", re.I),
    "pct": re.compile(r"%|процент|доля|pct|percent|share|удельн", re.I),
    "money": re.compile(r"сумм|руб|стоим|цена|выручк|затрат|тариф|amount|price|cost|revenue|₽", re.I),
}
YEAR_RE = re.compile(r"^(?P<pre>.*?)[\s_,]*(?P<year>(?:19|20)\d\d)(?!\d)(?P<post>.*)$")
SUBROW_RE = re.compile(r"^\s*из них\b", re.I)
TOTAL_RE = re.compile(r"^\s*(всего|итого|total)\b", re.I)
INT_TEXT = re.compile(r"^-?[1-9]\d{0,14}$")
DATE_LIKE = re.compile(r"^\d{1,4}[.\-/]\d{1,2}[.\-/]\d{1,4}([ T]\d{1,2}:\d{2}(:\d{2})?)?$")
NULLS = {"", "\\N", "NULL", "null", "None", "nan", "NaN", "#N/A"}


# ---------------------------------------------------------------- конфиг
def deep_merge(base, extra):
    out = copy.deepcopy(base)
    for k, v in (extra or {}).items():
        out[k] = deep_merge(out[k], v) if isinstance(v, dict) and isinstance(out.get(k), dict) else copy.deepcopy(v)
    return out


def load_config(path):
    if not path:
        return copy.deepcopy(DEFAULTS)
    with open(path, encoding="utf-8") as f:
        if path.lower().endswith(".json"):
            data = json.load(f)
        else:
            if yaml is None:
                sys.exit("Для конфига .yaml нужен pyyaml: pip install pyyaml (или используйте report.json)")
            data = yaml.safe_load(f) or {}
    return deep_merge(DEFAULTS, data)


def pattern_match(pattern, name):
    """Ключ правила: точное имя (без учёта регистра), маска (* ? [..]) или «re:регулярка»."""
    pattern = str(pattern)
    if pattern.startswith("re:"):
        return re.search(pattern[3:], name, re.I) is not None
    if any(ch in pattern for ch in "*?["):
        return fnmatch.fnmatch(name.lower(), pattern.lower())
    return pattern.lower() == name.lower()


def rules_for(rules, name):
    """Все подходящие правила по порядку записи сливаются; точное имя — сильнее масок."""
    merged = {}
    for pattern, rule in (rules or {}).items():
        if pattern_match(pattern, name) and any(ch in str(pattern) for ch in "*?[") | str(pattern).startswith("re:"):
            merged.update(rule or {})
    for pattern, rule in (rules or {}).items():
        if not str(pattern).startswith("re:") and not any(ch in str(pattern) for ch in "*?[") \
                and pattern_match(pattern, name):
            merged.update(rule or {})
    return merged


def file_config(cfg, file_name):
    """Настройки файла: общие + все подходящие маски из files (поздние перекрывают ранние)."""
    fc = {"table": cfg["table"], "coloring": cfg["coloring"], "columns": dict(cfg["columns"])}
    own = {}
    for pattern, rule in (cfg.get("files") or {}).items():
        if pattern_match(pattern, file_name) or pattern_match(pattern, os.path.splitext(file_name)[0]):
            own = deep_merge(own, rule or {})
    fc["table"] = deep_merge(cfg["table"], own.pop("table", {}))
    fc["coloring"] = deep_merge(cfg["coloring"], own.pop("coloring", {}))
    fc["columns"].update(own.pop("columns", {}) or {})
    fc.update(own)          # sheet, title, subtitle, description, order, skip
    return fc


# ---------------------------------------------------------------- чтение CSV
def read_csv_any(path):
    """-> (DataFrame строк, кодировка, разделитель). Все значения — строки, пустые — ''."""
    raw = open(path, "rb").read()
    text = None
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text, encoding = raw.decode(enc), enc
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("не удалось определить кодировку")
    if not text.strip():
        raise ValueError("пустой файл")
    sample = "\n".join(text.splitlines()[:50])
    try:
        sep = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        first = text.splitlines()[0]
        sep = max(";,\t|", key=first.count)
    from io import StringIO
    df = pd.read_csv(StringIO(text), sep=sep, dtype=str, keep_default_na=False, na_filter=False,
                     engine="python" if len(sep) > 1 else "c")
    df.columns = [fix_header(c) for c in df.columns]
    df = df.apply(lambda s: s.str.strip())
    df = df.mask(df.isin(NULLS), "")
    return df, encoding, sep


def fix_header(name):
    name = str(name).strip().strip('"')
    if name.startswith("?%") or name.startswith("? %"):
        return "Δ%" + name.split("%", 1)[1]
    if name.startswith("? "):
        return "Δ " + name[2:]
    return name


def to_numbers(series):
    """Строки -> float (NaN, где не число). Пробелы-разделители, десятичная запятая, «12,5%» (-> 0.125)."""
    s = series.str.replace(r"[\s\u00a0\u202f]", "", regex=True)
    pct = s.str.endswith("%")
    s = s.str.rstrip("%")
    both = s.str.contains(",", regex=False) & s.str.contains(".", regex=False)
    s = s.where(~both, s.str.replace(",", "", regex=False))        # 1,234.56 -> 1234.56
    s = s.str.replace(",", ".", regex=False)                        # 1234,56 -> 1234.56
    num = pd.to_numeric(s, errors="coerce")
    return num.where(~pct, num / 100), bool(pct[series != ""].all()) if (series != "").any() else False


# ---------------------------------------------------------------- колонки
class Col:
    def __init__(self, name):
        self.name = name
        self.label = name          # подпись в шапке (нижняя строка)
        self.group = None          # верхняя строка шапки
        self.role = "dimension"    # level | name | dimension | metric | service | subrow | skip
        self.kind = "text"         # text | date | datetime | count | money | number | pct | pct_points | delta | delta_pct
        self.fmt = None
        self.values = []           # значения для записи (None — пусто)
        self.formula = None        # ("diff"|"ratio", Col a, Col b) — добавленные колонки
        self.year = None
        self.base = None           # «Случаи» для «Случаи 2026»
        self.width = None
        self.hidden = False
        self.colorize = None
        self.positive_is = None
        self.description = ""
        self.cap = False           # Δ% > cap -> «> 200%»
        self.hint = None           # формула для справки у колонок из CSV (в ячейках — значения CSV)

    @property
    def is_metric(self):
        return self.role == "metric"

    @property
    def additive(self):
        return self.is_metric and self.kind in ("count", "money", "number", "delta")


def norm(text):
    return re.sub(r"[^0-9a-zа-яё]", "", str(text).lower())


def analyze_column(name, series, rule):
    col = Col(name)
    role, kind = rule.get("role"), rule.get("kind")
    nonempty = series[series != ""]
    if role is None:
        if RX["level"].match(name):
            role = "level"
        elif RX["name"].match(name):
            role = "name"
        elif RX["subrow"].match(name):
            role = "subrow"
    nums, all_pct_strings = to_numbers(series) if len(series) else (pd.Series(dtype=float), False)
    numeric_share = nums[series != ""].notna().mean() if len(nonempty) else 0.0
    leading_zero = nonempty.str.match(r"^0\d").any() if len(nonempty) else False
    if role is None:
        if len(nonempty) and numeric_share >= 0.98 and not leading_zero \
                and not RX["id"].search(name) and not RX["period"].match(name):
            role = "metric"
        else:
            role = "dimension"
    col.role = role

    if role == "metric":
        if kind is None:
            fractional = bool(((nums.dropna() % 1) != 0).any())
            if RX["delta_pct"].search(name):
                kind = "delta_pct"
            elif RX["delta"].match(name):
                kind = "delta"
            elif all_pct_strings or RX["pct"].search(name):
                kind = "pct" if (all_pct_strings or nums.abs().max() <= 1.5) else "pct_points"
            elif RX["money"].search(name) or fractional:
                kind = "money"
            else:
                kind = "count"
        col.kind = kind
        col.values = [None if pd.isna(v) else (int(v) if kind == "count" and float(v).is_integer() else float(v))
                      for v in nums]
        if kind == "delta":   # формат Δ — как у исходной метрики: целые или с копейками
            col.fmt = FMT["money"] if bool(((nums.dropna() % 1) != 0).any()) or RX["money"].search(name) else FMT["count"]
        col.cap = kind == "delta_pct"
    elif role in ("dimension", "service") and kind in (None, "date", "datetime") and len(nonempty) \
            and nonempty.str.match(DATE_LIKE).mean() >= 0.98:
        try:
            dates = pd.to_datetime(series.where(series != ""), dayfirst=True, errors="coerce", format="mixed")
        except (TypeError, ValueError):   # pandas < 2.0 — без format="mixed"
            dates = pd.to_datetime(series.where(series != ""), dayfirst=True, errors="coerce")
        if dates[series != ""].notna().all():
            has_time = bool((dates.dropna().dt.time != dt.time(0)).any())
            col.kind = "datetime" if has_time else "date"
            col.values = [None if pd.isna(v) else v.to_pydatetime() for v in dates]
    if not col.values:
        if role == "level" and len(nonempty) and nums[series != ""].notna().all() \
                and bool((nums.dropna() % 1 == 0).all()):
            col.kind = "count"
            col.values = [None if pd.isna(v) else int(v) for v in nums]
        else:
            if role == "level":           # «Уровень» не из целых чисел — обычный разрез
                col.role = role = "dimension"
            col.kind = kind or "text"
            # коды из цифр без ведущих нулей — числом (Excel не ругается на «число как текст»); kind: text — как есть
            keep_text = kind == "text"
            col.values = [None if v == "" else (int(v) if INT_TEXT.match(v) and not keep_text else v)
                          for v in series]

    col.fmt = rule.get("format") or col.fmt or FMT.get(col.kind)
    col.label = rule.get("label", name)
    col.group = rule.get("group")
    col.width = rule.get("width")
    col.hidden = bool(rule.get("hidden", False))
    col.colorize = rule.get("colorize")
    col.positive_is = rule.get("positive_is")
    col.description = rule.get("description", "")
    if rule.get("role") == "skip":
        col.role = "skip"
    return col


def apply_year_groups(cols, period, compare):
    """«X 2026» / «X 2025» -> группы шапки по годам; при compare=auto и отсутствии Δ — добавить Δ и Δ%."""
    metrics = [c for c in cols if c.is_metric]
    for c in metrics:
        if c.kind in ("delta", "delta_pct"):
            continue
        m = YEAR_RE.match(c.name)
        if m:
            c.year = int(m.group("year"))
            pre, post = m.group("pre").strip(" _,"), m.group("post").strip(" _")
            c.base = (pre + post if post.startswith(",") else f"{pre} {post}").strip(" _,")
    years = sorted({c.year for c in metrics if c.year}, reverse=True)
    if len(years) < 2:
        for c in metrics:
            c.year = None
        return cols
    y1, y0 = years[0], years[1]
    suffix = f" {period}" if period else ""
    by_base = {}
    for c in metrics:
        if c.year:
            by_base.setdefault(norm(c.base), {})[c.year] = c
            if c.group is None:
                c.group = f"{c.year}{suffix}"
            if c.label == c.name:
                c.label = c.base or c.name
    has_delta = any(c.kind == "delta" for c in metrics)
    pct_names = {norm(re.sub(RX["delta_pct"], "", c.name)) for c in metrics if c.kind == "delta_pct"}
    added_delta, added_pct = [], []
    for key, pair in by_base.items():
        if y1 not in pair or y0 not in pair or pair[y1].kind not in ("count", "money", "number"):
            continue
        a, b = pair[y1], pair[y0]
        if compare == "auto" and not has_delta:
            d = Col(f"Δ {a.base}")
            d.role, d.kind, d.formula, d.fmt = "metric", "delta", ("diff", a, b), a.fmt
            d.values = [None if (x is None and y is None) else round((x or 0) - (y or 0), 6)
                        for x, y in zip(a.values, b.values)]
            d.description = f"добавлено: {a.name} − {b.name}"
            added_delta.append(d)
        if compare == "auto" and key not in pct_names:
            p = Col(f"Δ% {a.base}")
            p.role, p.kind, p.formula, p.fmt, p.cap = "metric", "delta_pct", ("ratio", a, b), FMT["delta_pct"], True
            p.values = [None if not y else (x or 0) / y - 1 for x, y in zip(a.values, b.values)]
            p.description = f"добавлено: {a.name} / {b.name} − 1 (при {b.name} = 0 — пусто)"
            added_pct.append(p)
    for c in metrics:   # Δ из CSV: если нашлась пара по названию — формула для ручной правки (в ячейках — значения CSV)
        if c.kind == "delta":
            pair = by_base.get(norm(re.sub(RX["delta"], "", c.name, count=1)))
            if pair and y1 in pair and y0 in pair:
                c.hint = ("diff", pair[y1], pair[y0])
    cols = cols + added_delta + added_pct
    for c in cols:
        if c.is_metric and c.kind == "delta" and c.group is None:
            c.group = f"Δ {y1}-{y0}"
            c.label = re.sub(r"^(Δ|δ|дельта|delta|diff|разница|изм\w*)\s*", "", c.label, flags=re.I) or c.label
        elif c.is_metric and c.kind == "delta_pct" and c.group is None:
            c.group = f"Δ% {y1}/{y0}"
            c.label = re.sub(RX["delta_pct"], "", c.label).strip() or c.label
    # порядок метрик: годы по убыванию, Δ, Δ%, прочие метрики; остальные колонки — на своих местах
    rank = {y: i for i, y in enumerate(years)}

    def metric_order(c):
        if c.year:
            return (0, rank[c.year])
        return {"delta": (1, 0), "delta_pct": (2, 0)}.get(c.kind, (3, 0))
    metric_cols = sorted([c for c in cols if c.is_metric], key=metric_order)
    it = iter(metric_cols)
    first = next(i for i, c in enumerate(cols) if c.is_metric)
    rest = [c for c in cols if not c.is_metric]
    return rest[:first] + list(it) + rest[first:]


# ---------------------------------------------------------------- подготовка отчёта
class Report:
    pass


def prepare(path, file_name, fc, period):
    df, encoding, sep = read_csv_any(path)
    if df.empty:
        raise ValueError("нет строк (только шапка)")
    rules = fc["columns"]
    table = fc["table"]
    cols = []
    for name in df.columns:
        rule = rules_for(rules, name)
        c = analyze_column(name, df[name], rule)
        if c.role != "skip":
            cols.append(c)
    if not any(c.is_metric for c in cols):
        print(f"   [!] {file_name}: числовых колонок-метрик не найдено — лист без раскраски и итогов")
    level = next((c for c in cols if c.role == "level"), None)
    if level is not None and any(v is None for v in level.values):
        level.role = "dimension"
        level = None
    rep = Report()
    rep.file, rep.encoding, rep.sep = file_name, encoding, sep
    rep.hierarchy = level is not None
    rep.n_rows = len(df)
    if rep.n_rows > EXCEL_MAX_ROWS - 20:
        print(f"   [!] {file_name}: {rep.n_rows} строк — больше лимита Excel, лишние строки обрезаны")
        rep.n_rows = EXCEL_MAX_ROWS - 20
        for c in cols:
            c.values = c.values[:rep.n_rows]
    cols = apply_year_groups(cols, period, table["compare"])

    rows = list(range(rep.n_rows))
    if rep.hierarchy:
        rows = hierarchy_order(cols, level, table)
    elif table.get("sort"):
        rows = flat_order(cols, table["sort"], file_name)
    for c in cols:
        c.values = [c.values[i] for i in rows]

    # раскладка колонок
    if rep.hierarchy:
        name = next((c for c in cols if c.role == "name"), None) or next(
            (c for c in cols if c.role == "dimension" and c.kind == "text"), None)
        if name is None:
            raise ValueError("иерархический отчёт без колонки «Наименование»")
        name.role = "name"
        keys = [c for c in cols if c.role in ("dimension", "subrow") and c is not name]
        name.label = rules_for(rules, name.name).get("label") or " / ".join(level_labels(level, keys, rules)) or name.name
        for c in keys:
            c.role = "service" if c.role == "dimension" else c.role
        level.role = "service"
        metrics = [c for c in cols if c.is_metric]
        service = [level] + [c for c in cols if c.role in ("service", "subrow") and c is not level]
        rep.columns = [name] + metrics + service
        rep.level = level
        rep.subrow = next((c for c in cols if c.role == "subrow"), None)
        rep.name_col = name
    else:
        rep.columns = [c for c in cols if c.role != "level"] + [c for c in cols if c.role == "level"]
        rep.level, rep.subrow = None, None
        rep.name_col = next((c for c in rep.columns if not c.is_metric), None)
    rep.metrics = [c for c in rep.columns if c.is_metric]
    rep.table = table
    rn = table.get("row_numbers", "auto")
    rep.row_numbers = rep.hierarchy if rn == "auto" else bool(rn)
    return rep


def level_labels(level, keys, rules):
    """Подписи уровней: колонка разреза относится к уровню, на котором она впервые заполнена."""
    by_level = {}
    for c in keys:
        if c.role == "subrow":
            continue
        lv = [l for l, v in zip(level.values, c.values) if v not in (None, "")]
        if lv:
            by_level.setdefault(min(lv), []).append(c)
    labels = []
    for _, group in sorted(by_level.items()):
        names = []
        for c in group:
            label = rules_for(rules, c.name).get("level_label") or c.label
            if label not in names:
                names.append(label)
        labels.append(", ".join(names))
    return labels


def main_metric(cols, table):
    sort = table.get("sort")
    if sort:
        name = sort.lstrip("-+")
        found = next((c for c in cols if c.name.lower() == name.lower()), None)
        if found:
            return found, not sort.startswith("-")
    metrics = [c for c in cols if c.is_metric and c.kind in ("count", "money", "number")]
    if not metrics:
        return None, False
    newest = max((c.year or 0) for c in metrics)
    pool = [c for c in metrics if (c.year or 0) == newest]
    return next((c for c in pool if c.kind == "money"), pool[0]), False


def subrow_flags(cols, n):
    """Подстроки «из них …»: заполнена колонка «Из них» или наименование начинается с «из них»."""
    sub = next((c for c in cols if c.role == "subrow"), None)
    name = next((c for c in cols if c.role == "name"), None)
    return [bool(sub is not None and sub.values[i] not in (None, ""))
            or bool(name is not None and isinstance(name.values[i], str) and SUBROW_RE.match(name.values[i]))
            for i in range(n)]


def hierarchy_order(cols, level, table):
    """Структура — как в CSV: родитель строки — ближайшая строка выше с меньшим уровнем; подстроки «из них»
    прикреплены к строке над ними и ездят вместе с ней. Сортируются только дети внутри каждой группы
    (по главной метрике, при равенстве — по той же метрике прошлого года, затем порядок CSV)."""
    n = len(level.values)
    if str(table.get("hierarchy_sort", "auto")).lower() == "off":
        return list(range(n))
    metric, ascending = main_metric(cols, table)
    second = None
    if metric is not None and metric.year:
        second = next((c for c in cols if c.is_metric and c.base == metric.base and c.year and c.year < metric.year), None)
    is_sub = subrow_flags(cols, n)

    root = {"i": None, "level": -1, "children": [], "subs": []}
    stack, last = [root], root
    for i in range(n):
        if is_sub[i]:
            last["subs"].append(i)
            continue
        node = {"i": i, "level": level.values[i], "children": [], "subs": []}
        while len(stack) > 1 and stack[-1]["level"] >= node["level"]:
            stack.pop()
        stack[-1]["children"].append(node)
        stack.append(node)
        last = node

    def val(c, i):
        v = c.values[i] if c is not None else None
        return v if isinstance(v, (int, float)) else 0.0

    sign = 1 if ascending else -1

    def key(node):
        i = node["i"]
        return (sign * val(metric, i), sign * val(second, i), i)

    out = list(root["subs"])

    def walk(node):
        out.append(node["i"])
        out.extend(node["subs"])
        for child in sorted(node["children"], key=key):
            walk(child)
    for child in sorted(root["children"], key=key):
        walk(child)
    return out


def flat_order(cols, sort, file_name):
    name = sort.lstrip("-+")
    col = next((c for c in cols if c.name.lower() == name.lower()), None)
    if col is None:
        print(f"   [!] {file_name}: нет колонки для сортировки «{name}» — порядок CSV")
        return list(range(len(cols[0].values)))
    desc = sort.startswith("-")

    def key(i):
        v = col.values[i]
        return (v is None, (-v if desc and isinstance(v, (int, float)) else v) if v is not None else 0)
    if not col.is_metric:
        idx = sorted(range(len(col.values)), key=lambda i: ("" if col.values[i] is None else str(col.values[i])),
                     reverse=desc)
        return idx
    return sorted(range(len(col.values)), key=key)


# ---------------------------------------------------------------- оформление
class Styler:
    def __init__(self, style):
        self.s = style
        self._fonts, self._fills, self._align = {}, {}, {}

    def font(self, size=None, bold=False, color=C_BLACK, italic=False, underline=None):
        key = (size, bold, color, italic, underline)
        if key not in self._fonts:
            self._fonts[key] = Font(name=self.s["font"], size=size or self.s["size"], bold=bold, color=color,
                                    italic=italic, underline=underline)
        return self._fonts[key]

    def fill(self, color):
        if color not in self._fills:
            self._fills[color] = PatternFill("solid", fgColor=color)
        return self._fills[color]

    def align(self, h="center", indent=0, wrap=False):
        key = (h, indent, wrap)
        if key not in self._align:
            self._align[key] = Alignment(horizontal=h, vertical="center", indent=indent, wrap_text=wrap)
        return self._align[key]

    def head(self, ws, row, col, value=None, fill="F2F2F2", color=C_BLACK, size=None):
        cell = ws.cell(row, col)
        if value is not None:
            cell.value = value
        cell.font = self.font(size, True, color)
        cell.fill = self.fill(fill)
        cell.alignment = self.align("center", wrap=True)
        cell.border = HEAD_BORDER
        return cell


def link(cell, sheet, st, size=11, text=None):
    if text is not None:
        cell.value = text
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1")
    cell.font = st.font(size, False, LINK_COLOR, underline="single")


def mix(color, share):
    rgb = [int(color[k:k + 2], 16) for k in (0, 2, 4)]
    return "".join(f"{round(255 + (x - 255) * share):02X}" for x in rgb)


def colors_for(col, coloring):
    positive = (col.positive_is or coloring["positive_is"]).lower()
    return (CF_GREEN, CF_RED) if positive == "good" else (CF_RED, CF_GREEN)   # (рост, снижение)


def static_colors(values, col, coloring):
    """Мягкая шкала блока: 0 — белый, края — min/max блока (для Δ% — не выше pct_cap)."""
    nums = [v for v in values if isinstance(v, (int, float))]
    if not nums:
        return [None] * len(values)
    lo, hi = min(min(nums), 0.0), max(max(nums), 0.0)
    if col.kind in ("delta_pct", "pct"):
        hi = min(hi, coloring["pct_cap"])
    grow, fall = colors_for(col, coloring)
    gamma = coloring["gamma"]
    out = []
    for v in values:
        if not isinstance(v, (int, float)) or v == 0:
            out.append(None)
        elif v > 0 and hi > 0:
            out.append(mix(grow, min(v / hi, 1.0) ** gamma))
        elif v < 0 and lo < 0:
            out.append(mix(fall, min(v / lo, 1.0) ** gamma))
        else:
            out.append(None)
    return [None if c == CF_WHITE else c for c in out]


def colored_columns(rep, coloring):
    wanted = coloring.get("columns")
    out = []
    for c in rep.metrics:
        if c.colorize is not None:
            ok = bool(c.colorize)
        elif wanted:
            ok = any(pattern_match(p, c.name) for p in wanted)
        else:
            ok = c.kind in ("delta", "delta_pct")
        if ok:
            out.append(c)
    return out


def write_sheet(ws, rep, fc, cfg, st):
    style, table, coloring = cfg["style"], fc["table"], fc["coloring"]
    size = style["size"]
    cols = rep.columns
    offset = 1 if rep.row_numbers else 0
    ncols = len(cols) + offset
    col_idx = {id(c): offset + 1 + j for j, c in enumerate(cols)}
    letter = {id(c): get_column_letter(col_idx[id(c)]) for c in cols}
    n = rep.n_rows

    # --- строки 1–4: ссылка, заголовок, подзаголовок, описание
    link(ws.cell(1, 1), TOC_SHEET, st, 10, "← Содержание")
    title_col = 1 + offset
    ws.cell(2, title_col, fc.get("title") or rep.title).font = st.font(16, True, style["title_color"])
    subtitle = fc.get("subtitle", cfg["book"].get("subtitle"))
    if subtitle:
        ws.cell(3, title_col, subtitle).font = st.font(size, False, style["subtitle_color"])
    ws.cell(4, title_col, rep.description).font = st.font(size, False, style["muted_color"])
    for r in (2, 3, 4):
        ws.cell(r, title_col).alignment = st.align("left")

    totals_on = not rep.hierarchy and str(table.get("totals", "auto")).lower() not in ("off", "false", "0") \
        and any(c.additive for c in rep.metrics)
    first_col_text = rep.name_col.values if rep.name_col is not None else [None] * n
    is_total = [isinstance(v, str) and bool(TOTAL_RE.match(v)) for v in first_col_text] if not rep.hierarchy else [False] * n
    if str(table.get("totals", "auto")).lower() == "auto" and any(is_total):
        totals_on = False
    row = 5
    total_row = None
    if totals_on:
        total_row, row = row, row + 1
    two_level = any(c.group for c in cols)
    head1, head2 = row, row + (1 if two_level else 0)
    numbers_row = head2 + 1 if table.get("column_numbers", True) else None
    first_data = (numbers_row or head2) + 1
    last_data = first_data + n - 1

    # --- шапка
    if rep.row_numbers:
        for r in range(head1, head2 + 1):
            st.head(ws, r, 1, size=10)
        ws.cell(head1, 1).value = "№"
        if two_level:
            ws.merge_cells(start_row=head1, start_column=1, end_row=head2, end_column=1)
    j = 0
    while j < len(cols):
        c = cols[j]
        ci = col_idx[id(c)]
        kind = "service" if c.role in ("service", "subrow") else (c.kind if c.is_metric else "dimension")
        fill = HEAD_FILLS.get(kind, "F2F2F2")
        color = C_GREY if c.role in ("service", "subrow") else C_BLACK
        hsize = 11 if c.role in ("service", "subrow") else (size if len(str(c.label)) <= 60 else 12)
        if two_level and c.group:
            k = j
            while k + 1 < len(cols) and cols[k + 1].group == c.group:
                k += 1
            gfill = GROUP_FILL_DELTA if str(c.group).startswith("Δ") else GROUP_FILL_YEAR
            for m in range(j, k + 1):
                st.head(ws, head1, col_idx[id(cols[m])], fill=gfill)
                mk = cols[m].kind
                st.head(ws, head2, col_idx[id(cols[m])], cols[m].label, HEAD_FILLS.get(mk, "F2F2F2"),
                        C_GREY if mk in ("delta", "delta_pct") else C_BLACK)
            ws.cell(head1, ci).value = c.group
            if k > j:
                ws.merge_cells(start_row=head1, start_column=ci, end_row=head1, end_column=col_idx[id(cols[k])])
            j = k + 1
            continue
        for r in range(head1, head2 + 1):
            st.head(ws, r, ci, fill=fill, color=color, size=hsize)
        ws.cell(head1, ci).value = c.label
        if two_level and c.role not in ("service", "subrow"):   # служебные — без объединения, скрываются по одной
            ws.merge_cells(start_row=head1, start_column=ci, end_row=head2, end_column=ci)
        elif two_level:
            ws.cell(head1, ci).value = "служебная"
            ws.cell(head1, ci).font = st.font(9, True, C_GREY)
            ws.cell(head2, ci).value = c.label
        j += 1
    if numbers_row:
        for ci in range(1, ncols + 1):
            cell = ws.cell(numbers_row, ci, "+" if (ci == 1 and rep.row_numbers) else ci - offset)
            cell.font = st.font(10, False, C_GREY_LIGHT)
            cell.alignment = st.align("center")
            cell.border = HEAD_BORDER

    # --- данные
    levels = rep.level.values if rep.hierarchy else [None] * n
    is_sub = subrow_flags(cols, n) if rep.hierarchy else [False] * n
    if rep.hierarchy:
        next_main, nxt = [None] * n, None
        for i in range(n - 1, -1, -1):
            next_main[i] = nxt
            if not is_sub[i]:
                nxt = i
        is_leaf = [is_sub[i] or next_main[i] is None or levels[next_main[i]] <= levels[i] for i in range(n)]
    else:
        is_leaf = [not t for t in is_total]
    sub_font = st.font(size - 2, False, "7F7F7F", italic=True)
    service_font = st.font(11, False, C_SERVICE)
    no_font = st.font(10, False, C_GREY_LIGHT)
    pct_cap = coloring["pct_cap"]
    live = str(table.get("formulas", "values")).lower() == "live"
    tech = []          # (колонка, ячейки, формула/правило, пояснение) -> лист «Формулы» и примечания к заголовкам
    head_cell = {}     # колонка -> ячейка подписи в шапке (для примечания)
    for c in cols:
        head_cell[id(c)] = ws.cell(head2 if (two_level and c.group) else head1, col_idx[id(c)])

    for i in range(n):
        r = first_data + i
        lvl = levels[i]
        if rep.hierarchy:
            fill = LEVEL_FILLS.get(lvl, "EDEDED") if (not is_leaf[i] or lvl == 0) else None
            bold = lvl <= 1
        else:
            fill = TOTAL_FILL if is_total[i] else None
            bold = is_total[i]
        sub = is_sub[i]
        fill_obj = st.fill(fill) if fill else None
        if rep.row_numbers:
            cell = ws.cell(r, 1, i + 1)
            cell.font, cell.alignment, cell.border = no_font, st.align("center"), DATA_BORDER
        for c in cols:
            ci = col_idx[id(c)]
            value = formula_text(c.formula, r, letter) if (c.formula and live) else c.values[i]
            cell = ws.cell(r, ci, value)
            cell.border = DATA_BORDER
            if c.role in ("service", "subrow"):
                cell.font, cell.alignment = service_font, st.align("left")
            elif c.role == "name":
                cell.font = sub_font if sub else st.font(None, bold)
                cell.alignment = st.align("left", indent=min(lvl or 0, 15))
            elif c.is_metric:
                cell.font = sub_font if sub else st.font(None, bold)
                cell.alignment = st.align("center")
            else:
                cell.font = st.font(None, bold)
                cell.alignment = st.align("left" if c.kind == "text" else "center")
            if c.fmt:
                cell.number_format = c.fmt
            if c.cap and isinstance(c.values[i], (int, float)) and c.values[i] > pct_cap:
                cell.number_format = FMT_PCT_CAP
            if fill_obj is not None:
                cell.fill = fill_obj
        ws.row_dimensions[r].height = 18.75 * size / 14

    # --- строка «Итого» над шапкой (SUBTOTAL 109 — только видимые после фильтра строки)
    if total_row:
        label_col = col_idx[id(cols[0])]
        if rep.row_numbers:
            label_col = 1
        totals = {}
        for c in cols:
            if c.additive and not c.formula:
                totals[id(c)] = round(sum(v for v in c.values if isinstance(v, (int, float))), 6)
        for c in cols:
            if c.formula:
                op, a, b = c.formula
                x, y = totals.get(id(a)), totals.get(id(b))
                totals[id(c)] = (None if x is None or y is None else round(x - y, 6) if op == "diff"
                                 else (x / y - 1 if y else None))
        for c in cols:
            ci = col_idx[id(c)]
            cell = ws.cell(total_row, ci)
            L = letter[id(c)]
            if c.additive and not c.formula:
                text = f"=SUBTOTAL(109,{L}{first_data}:{L}{last_data})"
                cell.value = text if live else totals[id(c)]
                tech.append((c, f"{L}{total_row}", text, "строка «Итого»: сумма по видимым (после фильтра) строкам"))
            elif c.formula:
                text = formula_text(c.formula, total_row, letter)
                cell.value = text if live else totals[id(c)]
                tech.append((c, f"{L}{total_row}", text, "строка «Итого»"))
            over_cap = c.cap and not live and isinstance(totals.get(id(c)), (int, float)) and totals[id(c)] > pct_cap
            cell.number_format = FMT_PCT_CAP if over_cap else (c.fmt or "General")
            cell.font = st.font(None, True)
            cell.fill = st.fill(TOTAL_FILL)
            cell.border = DATA_BORDER
            cell.alignment = st.align("center")
        cell = ws.cell(total_row, label_col)
        if cell.value is None:
            cell.value = "Итого (по видимым строкам)" if live else "Итого"
            cell.alignment = st.align("left")
        if rep.row_numbers:
            cell.fill, cell.font, cell.border = st.fill(TOTAL_FILL), st.font(10, True), DATA_BORDER
        ws.row_dimensions[total_row].height = 18.75 * size / 14

    # --- раскраска Δ / Δ%
    n_colored = 0
    mode = str(coloring.get("mode", "static")).lower()
    targets = colored_columns(rep, coloring) if mode != "off" else []
    blocks = color_blocks(rep, levels, is_leaf, is_sub, n)
    # формулы колонок: добавленные Δ / Δ% и подсказки для Δ из CSV
    for c in cols:
        spec = c.formula or c.hint
        if spec and n:
            L = letter[id(c)]
            note = ("добавлено скриптом" if c.formula else "в ячейках — значения из CSV; формула — для проверки/правки")
            if c.cap:
                note += f"; формат «> 200%» для значений больше {pct_cap:.0%}: {FMT_PCT_CAP}"
            tech.append((c, f"{L}{first_data}:{L}{last_data}", formula_text(spec, first_data, letter),
                         note + f"; вставить в {L}{first_data} и протянуть вниз"))
    for c in targets:
        ci = col_idx[id(c)]
        grow, fall = colors_for(c, coloring)
        cells, n_ranges = ranges_text(letter[id(c)], first_data, blocks)
        rule = (f"Цветовая шкала (3 цвета): минимум — #{fall}, число 0 — #FFFFFF, максимум — #{grow}"
                + (f"; для Δ% максимум не выше {coloring['pct_cap']:.0%}" if c.kind in ("pct", "delta_pct") else ""))
        tech.append((c, cells, rule,
                     f"раскраска: {n_ranges} диапазон(ов), шкала отдельно для каждого"
                     + ("" if mode == "conditional" else
                        f"; в файле цвет записан заливкой, насыщенность = доля^{coloring['gamma']} "
                        "(средние значения бледнее, чем в шкале Excel)")))
        if mode == "conditional":
            grow, fall = colors_for(c, coloring)
            for block in blocks:
                for a, b in contiguous(block):
                    ref = f"{letter[id(c)]}{first_data + a}:{letter[id(c)]}{first_data + b}"
                    ws.conditional_formatting.add(ref, ColorScaleRule(
                        start_type="min", start_color=fall, mid_type="num", mid_value=0, mid_color=CF_WHITE,
                        end_type="max", end_color=grow))
                    n_colored += b - a + 1
        else:
            for block in blocks:
                fills = static_colors([c.values[i] for i in block], c, coloring)
                for i, color in zip(block, fills):
                    if color:
                        ws.cell(first_data + i, ci).fill = st.fill(color)
                        n_colored += 1

    rep.tech = tech
    if cfg["book"].get("header_notes", True):
        by_col = {}
        for c, cells, text, note in tech:
            by_col.setdefault(id(c), []).append(f"{cells}: {text}\n({note})")
        for c in cols:
            if id(c) in by_col:
                body = "\n\n".join(by_col[id(c)])
                head_cell[id(c)].comment = Comment(body, "report_builder", width=420,
                                                   height=min(400, 60 + 18 * body.count("\n") + len(body) // 3))

    # --- ширины
    widths = {}
    scale = size / 14
    for c in cols:
        ci = col_idx[id(c)]
        if c.width:
            widths[ci] = c.width
            continue
        sample = c.values[:3000]
        if c.is_metric:
            longest = max((len(shown(v, c)) for v in sample if isinstance(v, (int, float))), default=4)
            if total_row and c.additive:   # строка «Итого» — сумма длиннее любого значения
                longest = max(longest, len(shown(sum(abs(v) for v in c.values if isinstance(v, (int, float))), c)) + 1)
        elif c.kind in ("date", "datetime"):
            longest = 10 if c.kind == "date" else 16
        else:
            longest = max((len(str(v)) for v in sample if v is not None), default=4)
        if c.role == "name":
            longest = max((len(str(v)) * 1.0 + 2.5 * (lv or 0) for v, lv in zip(sample, levels[:3000])), default=40)
            top = table.get("name_max_width", 90)
            w = min(top, max(40, longest * CHAR_W * scale + 2))
        else:
            label_word = max((len(w) for w in str(c.label).split()), default=4)
            top = table.get("max_width", 45) if c.role != "service" else 24
            w = min(top, max(8, longest * CHAR_W * scale + 2.5, label_word * HEAD_CHAR_W * scale + 2))
        widths[ci] = round(w, 1)
    if rep.row_numbers:
        widths[1] = max(5, len(str(n)) * 1.2 + 2)
    for ci, w in widths.items():
        ws.column_dimensions[get_column_letter(ci)].width = w
    for c in cols:
        if c.hidden:
            ws.column_dimensions[letter[id(c)]].hidden = True
    # высота шапки — по самому длинному переносу
    lines = max((wrapped_lines(c.label, (widths[col_idx[id(c)]] - 1) / (HEAD_CHAR_W * scale)) for c in cols), default=1)
    lines = min(lines, 6)
    if two_level:
        ws.row_dimensions[head1].height = 21 * scale
        ws.row_dimensions[head2].height = max(21, 18 * lines) * scale
    else:
        ws.row_dimensions[head1].height = max(21, 18 * lines) * scale
    if numbers_row:
        ws.row_dimensions[numbers_row].height = 13.5

    # --- вид, фильтр, закрепление, печать
    ws.sheet_view.showGridLines = False
    if style.get("view") == "page_break":
        ws.sheet_view.view = "pageBreakPreview"
        ws.sheet_view.zoomScaleSheetLayoutView = style["zoom"]
    ws.sheet_view.zoomScale = style["zoom"]
    ws.sheet_view.zoomScaleNormal = style["zoom"]
    fcols = table.get("freeze_columns", "auto")
    if fcols == "auto":
        lead = 0
        for c in cols:
            if c.is_metric or c.role in ("service", "subrow"):
                break
            lead += 1
        fcols = lead if 0 < lead <= 3 and sum(widths[col_idx[id(c)]] for c in cols[:lead]) <= 100 else (1 if lead else 0)
    ws.freeze_panes = ws.cell(first_data, offset + 1 + int(fcols))
    filter_row = numbers_row or head2
    if not rep.hierarchy and table.get("autofilter", True) and n:
        ws.auto_filter.ref = f"A{filter_row}:{get_column_letter(ncols)}{last_data}"
    total_width = sum(w for ci, w in widths.items() if not ws.column_dimensions[get_column_letter(ci)].hidden)
    paper = style.get("paper", "auto")
    print_setup(ws, style, f"{head1}:{filter_row}", paper == "A3" or (paper == "auto" and total_width > 180))
    return n_colored, len(blocks)


def shown(v, c):
    """Как число будет выглядеть в ячейке — для ширины колонки."""
    if c.kind in ("pct", "delta_pct"):
        return "> 200%" if c.cap and v > 2 else f"+{v:.1%}"
    if c.kind == "pct_points":
        return f"{v:.1f}%"
    if c.kind == "money" or (c.fmt and "0.00" in c.fmt):
        return f"{v:,.2f}"
    return f"{v:,.0f}"


def wrapped_lines(text, per_line):
    """Сколько строк займёт подпись при переносе по словам."""
    per_line = max(int(per_line), 1)
    lines, cur = 1, 0
    for word in str(text).split():
        need = len(word) if cur == 0 else cur + 1 + len(word)
        if need <= per_line:
            cur = need
        else:
            lines += 1 if cur else 0
            cur = len(word)
    return lines


def formula_text(spec, r, letter):
    op, a, b = spec
    la, lb = letter[id(a)], letter[id(b)]
    return f"={la}{r}-{lb}{r}" if op == "diff" else f'=IF(N({lb}{r})=0,"",{la}{r}/{lb}{r}-1)'


def ranges_text(col_letter, first, block_list, limit=40):
    parts = [f"{col_letter}{first + a}" if a == b else f"{col_letter}{first + a}:{col_letter}{first + b}"
             for block in block_list for a, b in contiguous(block)]
    return "; ".join(parts[:limit]) + (f"; … ещё {len(parts) - limit}" if len(parts) > limit else ""), len(parts)


def color_blocks(rep, levels, is_leaf, is_sub, n):
    """Иерархия: подряд идущие строки-детали одного уровня — отдельная шкала; подстроки «из них» — своя шкала.
    Плоская таблица: все строки, кроме «Всего/Итого», — одна шкала."""
    if not rep.hierarchy:
        return [[i for i in range(n) if is_leaf[i]]]
    blocks, current = [], []
    for i in range(n + 1):
        if i < n and is_sub[i]:
            continue
        detail = i < n and is_leaf[i] and levels[i] > 0
        if current and (not detail or levels[i] != levels[current[0]]):
            blocks.append(current)
            current = []
        if detail:
            current.append(i)
    subs = [i for i in range(n) if is_sub[i]]
    if subs:
        blocks.append(subs)
    return [b for b in blocks if b]


def contiguous(block):
    out, start, prev = [], None, None
    for i in block:
        if start is None:
            start = prev = i
        elif i == prev + 1:
            prev = i
        else:
            out.append((start, prev))
            start = prev = i
    if start is not None:
        out.append((start, prev))
    return out


# ---------------------------------------------------------------- книга
def natural_key(text):
    return [int(p) if p.isdigit() else p.lower() for p in re.split(r"(\d+)", text)]


def report_number(stem):
    """-> (номер, имя без номера): «rep_x_1_1_202610021617» -> («1.1», «rep x»); «19.1 ЖК» -> («19.1», «ЖК»);
    без номера -> ('', имя). Хвост-штамп даты выгрузки DBeaver (_202610021617) отбрасывается."""
    stem = re.sub(r"_\d{8,}$", "", stem)
    m = re.match(r"^(\d+(?:[._]\d+)*)(?=\D|$)", stem)
    if m:
        return m.group(1).replace("_", "."), humanize(stem[m.end():]).strip(" .-")
    m = re.search(r"_(\d{1,3})(?:_(\d{1,2}))?$", stem)
    if m:
        return m.group(1) + (f".{m.group(2)}" if m.group(2) else ""), humanize(stem[:m.start()])
    return "", humanize(stem)


def sheet_title(name, used):
    base = re.sub(r"[\[\]:*?/\\]", " ", str(name)).strip().strip("'")[:31] or "Отчёт"
    title, k = base, 2
    while title.lower() in used:
        suffix = f" ({k})"
        title, k = base[:31 - len(suffix)] + suffix, k + 1
    used.add(title.lower())
    return title


def humanize(stem):
    stem = re.sub(r"_\d{8,}$", "", stem)
    return re.sub(r"[_]+", " ", stem).strip()


def discover(src_dir, cfg):
    files = [f for f in os.listdir(src_dir) if f.lower().endswith(".csv")]
    out = []
    for f in files:
        fc = file_config(cfg, f)
        if fc.get("skip"):
            print(f"[пропуск] {f}: skip в конфиге")
            continue
        stem = os.path.splitext(f)[0]
        no, _ = report_number(stem)
        order = fc.get("order")
        if order is not None:      # order: «1.4.1» — сразу после отчёта 1.4; «0» — в начало
            no = str(order)
        key = ((0, tuple(int(x) for x in re.findall(r"\d+", no)), natural_key(f)) if no
               else (1, (), natural_key(f)))
        out.append((key, f, fc))
    return [(f, fc) for _, f, fc in sorted(out, key=lambda x: x[0])]


def structure_text(rep):
    if rep.hierarchy:
        return "Всего → " + " → ".join(p.strip() for p in str(rep.name_col.label).split("/"))
    dims = [c.label for c in rep.columns if not c.is_metric and c.role not in ("service", "subrow", "level")]
    return ", ".join(dims) or "—"


def key_total(rep, table):
    metric, _ = main_metric(rep.metrics, table)
    if metric is None:
        return None, None
    if rep.hierarchy:
        zero = [v for v, lv in zip(metric.values, rep.level.values) if lv == 0]
        return metric, (zero[0] if zero else None)
    name_vals = rep.name_col.values if rep.name_col is not None else [None] * rep.n_rows
    totals = [v for v, nv in zip(metric.values, name_vals) if isinstance(nv, str) and TOTAL_RE.match(nv)]
    if totals:
        return metric, totals[0]
    return metric, sum(v for v in metric.values if isinstance(v, (int, float)))


def build(src_dir, cfg, out_path, explain=False):
    st = Styler(cfg["style"])
    book = cfg["book"]
    files = discover(src_dir, cfg)
    if not files:
        sys.exit(f"В папке {src_dir} нет CSV-файлов")
    period = book.get("period") or ""
    wb = Workbook()
    toc = wb.active
    toc.title = TOC_SHEET
    used = {TOC_SHEET.lower(), HELP_SHEET.lower(), FORMULAS_SHEET.lower()}
    entries, skipped, help_rows = [], [], {}

    for f, fc in files:
        path = os.path.join(src_dir, f)
        stem = os.path.splitext(f)[0]
        try:
            rep = prepare(path, f, fc, period)
        except Exception as err:   # плохой файл не должен ронять сборку остальных
            print(f"[пропуск] {f}: {err}")
            skipped.append((f, str(err)))
            continue
        no, rest = report_number(stem)
        rep.title = fc.get("title") or book.get("title") or os.path.basename(src_dir.rstrip("\\/"))
        rep.description = fc.get("description") or (f"Отчёт {no}. {rest}" if no else rest)
        if fc.get("description") and no and not fc.get("description").startswith("Отчёт"):
            rep.description = f"Отчёт {no}. {fc['description']}"
        roles = ", ".join(f"{c.name}={c.role}{'/' + c.kind if c.is_metric else ''}{'(+)' if c.formula else ''}"
                          for c in rep.columns)
        if explain:
            print(f"[{'иерархия' if rep.hierarchy else 'таблица '}] {f}  ({rep.encoding}, '{rep.sep}', строк {rep.n_rows})")
            print(f"     {roles}")
            continue
        name = sheet_title(fc.get("sheet") or (f"{no} {rest}" if no else rest), used)
        ws = wb.create_sheet(name)
        n_colored, n_blocks = write_sheet(ws, rep, fc, cfg, st)
        print(f"[ok] {name:31} строк={rep.n_rows:7}  метрик={len(rep.metrics):3}  окрашено={n_colored:7}  <- {f}")
        entries.append((name, rep))
        for c in rep.columns:
            h = help_rows.setdefault(c.name, {"label": c.label, "role": c.role, "kind": c.kind, "fmt": c.fmt,
                                              "description": c.description, "sheets": []})
            h["sheets"].append(name)
    if explain:
        return
    if not entries:
        sys.exit("Ни один CSV не подошёл — книга не сохранена")

    write_toc(toc, entries, skipped, cfg, src_dir, st)
    if book.get("columns_sheet", True):
        write_help(wb.create_sheet(HELP_SHEET), help_rows, st, cfg["style"])
    mode = str(book.get("formulas_sheet", "hidden")).lower()
    if mode != "off" and any(rep.tech for _, rep in entries):
        fws = wb.create_sheet(FORMULAS_SHEET)
        write_formulas(fws, entries, st, cfg["style"])
        if mode == "hidden":
            fws.sheet_state = "hidden"
    wb.save(out_path)
    print(f"Готово: {out_path}")


def write_toc(ws, entries, skipped, cfg, src_dir, st):
    book, style = cfg["book"], cfg["style"]
    ws.sheet_view.showGridLines = False
    title = book.get("title") or os.path.basename(src_dir.rstrip("\\/"))
    ws["A1"] = title
    ws["A1"].font = st.font(16, True, style["title_color"])
    row = 2
    if book.get("subtitle"):
        ws.cell(row, 1, book["subtitle"]).font = st.font(12, False, style["subtitle_color"])
        row += 1
    ws.cell(row, 1, f"Источник: {src_dir}. Собрано {dt.datetime.now():%d.%m.%Y %H:%M}.").font = st.font(11, False, C_GREY)
    row += 1
    if str(book.get("formulas_sheet", "hidden")).lower() != "off":
        hidden = str(book.get("formulas_sheet", "hidden")).lower() == "hidden"
        ws.cell(row, 1, "В ячейках — значения (без формул и условного форматирования). Формулы Δ / Δ% / Итого и правила "
                        "раскраски — в примечаниях к заголовкам колонок и на листе «Формулы»"
                        + (" (скрыт: ПКМ по ярлыку листа → Показать)." if hidden else ".")).font = st.font(11, False, C_GREY)
        row += 1
    for note in book.get("notes") or []:
        ws.cell(row, 1, note).font = st.font(11, False, C_GREY)
        row += 1
    row += 1
    heads = ["№", "Лист", "Описание", "Структура", "Метрики", "Строк", "Ключевой итог", "Значение", "Файл"]
    for ci, h in enumerate(heads, start=1):
        st.head(ws, row, ci, h, size=11)
    head_row = row
    for k, (name, rep) in enumerate(entries, start=1):
        row += 1
        ws.cell(row, 1, k)
        link(ws.cell(row, 2), name, st, 11, name)
        ws.cell(row, 3, rep.description)
        ws.cell(row, 4, structure_text(rep))
        names = [c.name for c in rep.metrics if not c.formula]
        text = ", ".join(names)
        ws.cell(row, 5, (text if len(text) <= 120 else text[:117].rsplit(",", 1)[0] + f" … всего {len(names)}") or "—")
        ws.cell(row, 6, rep.n_rows).number_format = "#,##0"
        metric, value = key_total(rep, rep.table)
        ws.cell(row, 7, metric.name if metric else None)
        if value is not None:
            ws.cell(row, 8, value).number_format = metric.fmt or FMT["money"]
        ws.cell(row, 9, rep.file)
        for ci in range(1, len(heads) + 1):
            cell = ws.cell(row, ci)
            cell.border = DATA_BORDER
            if ci != 2:
                cell.font = st.font(11)
            cell.alignment = st.align("center" if ci in (1, 6, 8) else "left", wrap=ci in (3, 4, 5))
    if skipped:
        row += 2
        ws.cell(row, 1, "Пропущенные файлы").font = st.font(11, True, "C00000")
        for f, reason in skipped:
            row += 1
            ws.cell(row, 2, f).font = st.font(11)
            ws.cell(row, 3, reason).font = st.font(11, False, C_GREY)
    for col, width in zip("ABCDEFGHI", (5, 32, 60, 45, 45, 9, 24, 16, 40)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    print_setup(ws, cfg["style"], f"{head_row}:{head_row}")


def print_setup(ws, style, title_rows=None, a3=False):
    if title_rows:
        ws.print_title_rows = title_rows
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if a3 else ws.PAPERSIZE_A4
    ws.page_setup.orientation = style.get("orientation", "landscape")
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = 0.12
    ws.page_margins.bottom, ws.page_margins.header, ws.page_margins.footer = 0.4, 0.1, 0.15
    ws.oddFooter.center.text = "Страница &P из &N"
    ws.oddFooter.center.size = 10
    ws.oddFooter.center.font = f"{style['font']},Regular"
    ws.oddFooter.center.color = C_GREY


def write_formulas(ws, entries, st, style):
    """Формулы и правила раскраски текстом: скопировать из колонки «Формула / правило» в первую ячейку диапазона
    и протянуть вниз (ссылки относительные)."""
    ws.sheet_view.showGridLines = False
    link(ws.cell(1, 1), TOC_SHEET, st, 10, "← Содержание")
    ws.cell(2, 1, "Формулы и правила раскраски").font = st.font(16, True, "002060")
    ws.cell(3, 1, "В листах отчёта — только значения. Формулу из колонки D вставьте в первую ячейку диапазона и "
                  "протяните вниз; правило раскраски — Главная → Условное форматирование → Цветовые шкалы → "
                  "Другие правила.").font = st.font(11, False, C_GREY)
    heads = ["Лист", "Колонка", "Ячейки", "Формула / правило", "Пояснение"]
    for ci, h in enumerate(heads, start=1):
        st.head(ws, 5, ci, h, size=11)
    r = 5
    for name, rep in entries:
        for c, cells, text, note in rep.tech:
            r += 1
            link(ws.cell(r, 1), name, st, 11, name)
            ws.cell(r, 2, c.name)
            ws.cell(r, 3, cells)
            cell = ws.cell(r, 4, text)
            cell.data_type, cell.quotePrefix = "s", True     # текст «=…», Excel его не вычисляет
            ws.cell(r, 5, note)
            for ci in range(1, 6):
                cell = ws.cell(r, ci)
                cell.border = DATA_BORDER
                if ci != 1:
                    cell.font = st.font(11, False, "1F3864" if ci == 4 else C_BLACK)
                cell.alignment = st.align("left", wrap=ci in (3, 4, 5))
    for col, width in zip("ABCDE", (30, 30, 30, 60, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:E{r}"
    print_setup(ws, style, "5:5")


def write_help(ws, rows, st, style):
    ws.sheet_view.showGridLines = False
    link(ws.cell(1, 1), TOC_SHEET, st, 10, "← Содержание")
    ws.cell(2, 1, "Справка по колонкам").font = st.font(16, True, "002060")
    roles = {"metric": "метрика", "dimension": "разрез", "name": "наименование", "service": "служебная",
             "level": "уровень", "subrow": "подстрока «из них»"}
    heads = ["Колонка CSV", "Подпись в шапке", "Роль", "Вид", "Описание", "Листы"]
    for ci, h in enumerate(heads, start=1):
        st.head(ws, 4, ci, h, size=11)
    r = 4
    for name, h in rows.items():
        r += 1
        values = [name, h["label"], roles.get(h["role"], h["role"]), h["kind"], h["description"], ", ".join(h["sheets"])]
        for ci, v in enumerate(values, start=1):
            cell = ws.cell(r, ci, v)
            cell.font, cell.border = st.font(11), DATA_BORDER
            cell.alignment = st.align("left", wrap=ci in (5, 6))
    for col, width in zip("ABCDEF", (34, 30, 18, 12, 50, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{r}"
    print_setup(ws, style, "4:4")


# ---------------------------------------------------------------- шаблон конфига
def yq(text):
    return json.dumps(str(text), ensure_ascii=False)


def write_template(src_dir, cfg, path):
    lines = [
        "# Настройки сборки отчёта. Все поля необязательны: удалите лишнее — сработает автоопределение.",
        "# Маски имён: точное имя, * и ? (glob) или re:регулярка. Регистр не важен.",
        "",
        "book:",
        f"  title: {yq(os.path.basename(src_dir.rstrip(chr(92) + '/')))}     # заголовок книги и листов",
        '  subtitle: ""            # красная строка под заголовком (отбор, источник): «ЗЛ, ДЗМ, без МРВСМ»',
        '  period: ""              # подпись к годам в шапке: «9м» -> «2026 9м»',
        "  notes:                  # пояснения на листе «Содержание»",
        '    - "Суммы — млн руб."',
        "  # output: D:\\\\отчёты\\\\итог.xlsx",
        "",
        "style:",
        "  font: Arial Narrow",
        "  size: 14",
        "  view: page_break        # page_break | normal",
        "  zoom: 70",
        "  paper: auto             # auto | A3 | A4",
        "",
        "table:",
        "  compare: auto           # auto — добавить Δ и Δ% для пар «X 2026 / X 2025» | off",
        "  formulas: values        # values — в ячейках числа, формулы текстом на листе «Формулы» | live",
        "  totals: auto            # плоские таблицы: строка «Итого» над шапкой (по фильтру) | off",
        "  autofilter: true",
        "  # sort: \"-Сумма 2026\"   # минус — по убыванию",
        "  hierarchy_sort: auto    # иерархия (колонка «Уровень»): дети по убыванию главной метрики | off",
        "",
        "coloring:",
        "  mode: static            # static — заливкой (быстро, не «плывёт») | conditional — условное форм. | off",
        "  positive_is: bad        # bad — рост красный (затраты) | good — рост зелёный (выручка)",
        "  # columns: [\"Δ Случаи\", \"Δ% *\"]   # что красить; по умолчанию — все Δ и Δ%",
        "  pct_cap: 2.0            # Δ% > 200% -> «> 200%»",
        "",
        "# Правила для колонок во всех файлах. Ключи: role (metric | dimension | name | level | service | skip),",
        "# kind (count | money | number | pct | pct_points | delta | delta_pct | text | date), label, group,",
        "# format (формат Excel), width, hidden, colorize (true/false), positive_is, level_label, description.",
        "columns:",
        '  "re:^(id|gid|pat_id)$": {role: dimension}',
        '  # "Сумма*": {kind: money, description: "Сумма оплаты, млн руб."}',
        '  # "OTD_NAME": {label: "Отделение", level_label: "Отделение"}',
        "",
        "# Настройки файлов (по маске имени). Ключи: sheet (имя листа ≤31), title, subtitle, description,",
        "# order (место листа: «1.4.1» — сразу после 1.4), skip, а также table / coloring / columns — только для этого файла.",
        "files:",
    ]
    for f, fc in discover(src_dir, cfg):
        stem = os.path.splitext(f)[0]
        lines.append(f"  {yq(f)}:")
        lines.append(f"    sheet: {yq(humanize(stem)[:31])}")
        lines.append('    description: ""')
        try:
            rep = prepare(os.path.join(src_dir, f), f, fc, cfg["book"].get("period") or "")
            kind = "иерархия по «Уровень»" if rep.hierarchy else "плоская таблица"
            lines.append(f"    # {kind}; строк {rep.n_rows}; {rep.encoding}, разделитель {rep.sep!r}")
            lines.append("    # метрики: " + (", ".join(f"{c.name} [{c.kind}{', добавлено' if c.formula else ''}]"
                                                for c in rep.metrics) or "нет"))
            dims = [c.name for c in rep.columns if not c.is_metric]
            lines.append("    # разрезы: " + (", ".join(dims) or "нет"))
        except Exception as err:
            lines.append(f"    # не прочитан: {err}")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    print(f"Шаблон настроек: {path}")


# ---------------------------------------------------------------- запуск
def main():
    p = argparse.ArgumentParser(description="Excel-отчёт в едином оформлении из всех CSV в папке")
    p.add_argument("src", help="папка с CSV")
    p.add_argument("--out", help="итоговый .xlsx (по умолчанию <папка>.xlsx рядом с папкой)")
    p.add_argument("--config", help="конфиг .yaml/.json (по умолчанию report.yaml в папке, если есть)")
    p.add_argument("--init", action="store_true", help="создать шаблон report.yaml по найденным файлам")
    p.add_argument("--explain", action="store_true", help="показать распознанные колонки, без сборки")
    args = p.parse_args()
    src = os.path.abspath(args.src)
    if not os.path.isdir(src):
        sys.exit(f"Нет папки {src}")
    config_path = args.config or next((os.path.join(src, n) for n in CONFIG_NAMES
                                       if os.path.exists(os.path.join(src, n))), None)
    cfg = load_config(config_path)
    if config_path:
        print(f"Конфиг: {config_path}")
    if args.init:
        target = os.path.join(src, "report.yaml")
        if os.path.exists(target):
            target = os.path.join(src, "report.template.yaml")
        write_template(src, cfg, target)
        return
    out = args.out or cfg["book"].get("output") or src.rstrip("\\/") + ".xlsx"
    if os.path.exists(out) and not args.explain:
        try:
            os.rename(out, out)   # открыт в Excel -> на Windows ошибка
        except OSError:
            sys.exit(f"Файл {out} открыт в Excel — закройте его или укажите другой --out")
    build(src, cfg, out, explain=args.explain)


if __name__ == "__main__":
    main()
